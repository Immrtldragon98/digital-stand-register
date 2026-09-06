import io
import re
from datetime import date, datetime
from statistics import mean

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.auth.dependencies import require_operator
from app.database.session import get_db
from app.models.reliability_history import HistoricalCampaign, HistoricalSpareUsage, ProcessObservation
from app.models.user import User

router = APIRouter()

PARAMETER_SHEET_MAP = {
    "WRM 3 PARAMETER": (3, "W1"),
    "WRM 4 PARAMETER": (4, "W2"),
    "WRM 5 PARAMETER": (5, "W3"),
}


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _clean_stand(value):
    if value is None:
        return None
    text = str(value).strip().upper()
    text = re.sub(r"\s*\(PZ\)\s*", "", text).strip()
    if "/" in text:
        return None
    if re.fullmatch(r"\d+\.\d+", text):
        return text
    match = re.match(r"^0*(\d+)([A-Z]?)$", text)
    if not match:
        return text[:50] or None
    return f"{int(match.group(1))}{match.group(2)}"


def _num(value):
    try:
        if value is None or value == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _parse_stand_tracking(wb, filename):
    campaigns = []
    sheet_map = {"WRM1": "W1", "WRM2": "W2", "WRM3": "W3"}
    for sheet_name, line_name in sheet_map.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        active = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_date = _as_date(row[0] if len(row) else None)
            if not row_date:
                continue
            remark = str(row[12]).strip() if len(row) > 12 and row[12] not in (None, "") else None
            for position in range(1, 11):
                idx = position + 1
                stand = _clean_stand(row[idx] if idx < len(row) else None)
                if not stand:
                    continue
                current = active.get(position)
                if current is None:
                    active[position] = (stand, row_date)
                    continue
                previous_stand, started = current
                if previous_stand == stand:
                    continue
                life_days = max(0, (row_date - started).days)
                campaigns.append({
                    "line_name": line_name,
                    "position_number": position,
                    "stand_code": previous_stand,
                    "installed_date": started,
                    "removed_date": row_date,
                    "life_days": float(life_days),
                    "removal_reason": remark,
                    "inferred": True,
                    "confidence": 0.75 if remark else 0.65,
                    "source_file": filename,
                })
                active[position] = (stand, row_date)
        # Keep open campaigns out of the reliability average; live DSR now tracks those exactly.
    return campaigns


def _parse_spare_usage(wb, filename):
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    usages = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_date = _as_date(row[0] if row else None)
        if not row_date:
            continue
        for idx in range(1, min(len(row), len(headers))):
            spare = headers[idx].strip()
            qty = _num(row[idx])
            if not spare or qty is None or qty == 0:
                continue
            usages.append({"usage_date": row_date, "spare_name": spare[:150], "quantity": qty, "source_file": filename})
    return usages


def _parse_parameter_report(wb, filename):
    rows = []
    for sheet_name, (source_wrm, line_name) in PARAMETER_SHEET_MAP.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        for raw in ws.iter_rows(min_row=2, values_only=True):
            obs_date = _as_date(raw[0] if raw else None)
            if not obs_date:
                continue
            params = {}
            for i, key in enumerate(header):
                if key and i < len(raw):
                    value = raw[i]
                    if isinstance(value, (datetime, date)):
                        value = value.isoformat()
                    params[key] = value
            rows.append({
                "observation_date": obs_date,
                "shift": str(raw[1]).strip()[:20] if len(raw) > 1 and raw[1] is not None else None,
                "source_wrm": source_wrm,
                "line_name": line_name,
                "coil_no": str(raw[3]).strip()[:80] if len(raw) > 3 and raw[3] is not None else None,
                "uts_band": str(raw[4]).strip()[:30] if len(raw) > 4 and raw[4] is not None else None,
                "casting_speed": _num(raw[5] if len(raw) > 5 else None),
                "emulsion_temp": _num(raw[15] if len(raw) > 15 else None),
                "parameters": params,
                "source_file": filename,
            })
    return rows


def _detect(wb):
    names = set(wb.sheetnames)
    if {"WRM1", "WRM2", "WRM3"}.intersection(names):
        return "STAND_TRACKING"
    if any(name in names for name in PARAMETER_SHEET_MAP):
        return "PROCESS_PARAMETERS"
    first = wb[wb.sheetnames[0]]
    first_row = [str(c.value or "").strip().lower() for c in first[1]]
    if "oil seal" in first_row or "special bearing" in first_row:
        return "SPARE_USAGE"
    return "UNKNOWN"


@router.post("/preview")
async def preview_historical(file: UploadFile = File(...), _: User = Depends(require_operator)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Upload an .xlsx workbook")
    payload = await file.read()
    try:
        wb = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Could not read workbook: {exc}")
    kind = _detect(wb)
    if kind == "STAND_TRACKING":
        rows = _parse_stand_tracking(wb, file.filename)
        reasons = sum(1 for r in rows if r["removal_reason"])
        return {"kind": kind, "filename": file.filename, "records": len(rows), "with_reasons": reasons, "lines": sorted(set(r["line_name"] for r in rows))}
    if kind == "SPARE_USAGE":
        rows = _parse_spare_usage(wb, file.filename)
        return {"kind": kind, "filename": file.filename, "records": len(rows), "spares": len(set(r["spare_name"] for r in rows))}
    if kind == "PROCESS_PARAMETERS":
        rows = _parse_parameter_report(wb, file.filename)
        return {"kind": kind, "filename": file.filename, "records": len(rows), "mapping": {"WRM3": "W1", "WRM4": "W2", "WRM5": "W3"}}
    raise HTTPException(400, "Workbook format not recognized")


@router.post("/import")
async def import_historical(file: UploadFile = File(...), db: Session = Depends(get_db), _: User = Depends(require_operator)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Upload an .xlsx workbook")
    payload = await file.read()
    try:
        wb = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Could not read workbook: {exc}")
    kind = _detect(wb)
    inserted = 0
    skipped = 0
    if kind == "STAND_TRACKING":
        for row in _parse_stand_tracking(wb, file.filename):
            exists = db.query(HistoricalCampaign.id).filter_by(
                source_file=row["source_file"], line_name=row["line_name"], position_number=row["position_number"],
                stand_code=row["stand_code"], installed_date=row["installed_date"]
            ).first()
            if exists:
                skipped += 1
                continue
            db.add(HistoricalCampaign(**row)); inserted += 1
    elif kind == "SPARE_USAGE":
        for row in _parse_spare_usage(wb, file.filename):
            exists = db.query(HistoricalSpareUsage.id).filter_by(source_file=row["source_file"], usage_date=row["usage_date"], spare_name=row["spare_name"]).first()
            if exists:
                skipped += 1
                continue
            db.add(HistoricalSpareUsage(**row)); inserted += 1
    elif kind == "PROCESS_PARAMETERS":
        for row in _parse_parameter_report(wb, file.filename):
            exists = db.query(ProcessObservation.id).filter_by(
                source_file=row["source_file"], line_name=row["line_name"], observation_date=row["observation_date"],
                shift=row["shift"], coil_no=row["coil_no"]
            ).first()
            if exists:
                skipped += 1
                continue
            db.add(ProcessObservation(**row)); inserted += 1
    else:
        raise HTTPException(400, "Workbook format not recognized")
    db.commit()
    return {"kind": kind, "inserted": inserted, "skipped": skipped, "filename": file.filename}


@router.get("/summary")
def historical_summary(db: Session = Depends(get_db)):
    campaigns = db.query(HistoricalCampaign).all()
    by_position = []
    for line_name in ("W1", "W2", "W3"):
        for position in range(1, 11):
            values = [c.life_days for c in campaigns if c.line_name == line_name and c.position_number == position and c.life_days is not None and c.life_days >= 0]
            if values:
                by_position.append({"line": line_name, "position": position, "campaigns": len(values), "avg_days": round(mean(values), 2)})
    return {
        "campaigns": len(campaigns),
        "spare_usage_rows": db.query(HistoricalSpareUsage).count(),
        "process_observations": db.query(ProcessObservation).count(),
        "by_position": by_position,
    }
