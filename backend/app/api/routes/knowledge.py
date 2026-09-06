import csv
import hashlib
import io
import re
from pathlib import Path

from docx import Document as DocxDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.auth.dependencies import require_admin, require_operator
from app.database.session import get_db
from app.models.knowledge_document import KnowledgeChunk, KnowledgeDocument
from app.models.user import User

router = APIRouter()

ALLOWED_CATEGORIES = {"ROD_QUALITY", "EMULSION", "SPARE_USAGE", "FMEA_RCA", "SHIFT_REPORT", "OTHER"}
ALLOWED_SUFFIXES = {".txt", ".md", ".csv", ".xlsx", ".pdf", ".docx"}
MAX_BYTES = 12 * 1024 * 1024


def _extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise HTTPException(400, f"Unsupported file type {suffix}. Use TXT, MD, CSV, XLSX, PDF or DOCX.")

    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="ignore")

    if suffix == ".csv":
        text = data.decode("utf-8", errors="ignore")
        rows = csv.reader(io.StringIO(text))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)

    if suffix == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"SHEET: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                values = [str(v).strip() if v is not None else "" for v in row]
                if any(values):
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if suffix == ".docx":
        doc = DocxDocument(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)

    return ""


def _chunks(text: str, size: int = 1600, overlap: int = 250):
    clean = re.sub(r"\r\n?", "\n", text)
    clean = re.sub(r"[ \t]+", " ", clean).strip()
    if not clean:
        return []
    out = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + size)
        out.append(clean[start:end])
        if end == len(clean):
            break
        start = max(start + 1, end - overlap)
    return out


@router.get("/")
def list_documents(db: Session = Depends(get_db)):
    docs = db.query(KnowledgeDocument).order_by(KnowledgeDocument.created_at.desc()).all()
    return [{"id": d.id, "filename": d.filename, "category": d.category, "uploaded_by": d.uploaded_by, "created_at": d.created_at, "chunks": len(d.chunks)} for d in docs]


@router.post("/upload", status_code=201)
async def upload_document(
    category: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_operator),
):
    category = category.strip().upper()
    if category not in ALLOWED_CATEGORIES:
        raise HTTPException(400, "Unknown document category")
    data = await file.read()
    if not data:
        raise HTTPException(400, "File is empty")
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File is larger than 12 MB")

    digest = hashlib.sha256(data).hexdigest()
    existing = db.query(KnowledgeDocument).filter(KnowledgeDocument.content_hash == digest).first()
    if existing:
        raise HTTPException(409, f"This document is already imported as {existing.filename}")

    text = _extract_text(file.filename or "document.txt", data)
    chunks = _chunks(text)
    if not chunks:
        raise HTTPException(400, "No readable text was found in the document")

    doc = KnowledgeDocument(filename=file.filename or "document", category=category, content_hash=digest, uploaded_by=user.username)
    db.add(doc)
    db.flush()
    for idx, content in enumerate(chunks):
        db.add(KnowledgeChunk(document_id=doc.id, chunk_index=idx, content=content))
    db.commit()
    db.refresh(doc)
    return {"id": doc.id, "filename": doc.filename, "category": doc.category, "chunks": len(chunks), "message": "Document indexed for retrieval"}


@router.get("/search")
def search_knowledge(q: str, limit: int = 8, db: Session = Depends(get_db), _: User = Depends(require_operator)):
    terms = [t.lower() for t in re.findall(r"[A-Za-z0-9_.-]+", q) if len(t) > 2]
    if not terms:
        return []
    chunks = db.query(KnowledgeChunk).join(KnowledgeDocument).all()
    ranked = []
    for chunk in chunks:
        lower = chunk.content.lower()
        score = sum(lower.count(term) for term in terms)
        if score:
            ranked.append((score, chunk))
    ranked.sort(key=lambda x: x[0], reverse=True)
    return [{"score": score, "document_id": c.document_id, "filename": c.document.filename, "category": c.document.category, "chunk_index": c.chunk_index, "content": c.content} for score, c in ranked[:max(1, min(limit, 20))]]


@router.delete("/{document_id}")
def delete_document(document_id: int, db: Session = Depends(get_db), _: User = Depends(require_admin)):
    doc = db.get(KnowledgeDocument, document_id)
    if not doc:
        raise HTTPException(404, "Document not found")
    db.delete(doc)
    db.commit()
    return {"status": "deleted", "id": document_id}
